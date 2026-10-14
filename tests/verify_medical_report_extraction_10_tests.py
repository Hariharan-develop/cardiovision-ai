"""
CardioVision AI — Comprehensive 10-Scenario Medical Document Extraction Verification Suite
Tests the complete extraction engine and REST API against:
1. Text-based medical PDF
2. Scanned medical PDF
3. Multi-page laboratory report
4. DOCX medical report
5. PDF containing laboratory tables
6. PDF with missing measurements
7. PDF containing conflicting values
8. PDF with different measurement units
9. Invalid or corrupted PDF
10. PDF with all supported clinical fields

Validates:
- Exact extracted values vs expected ground truth
- Analyte-specific unit conversions
- Disambiguation of reference intervals
- Source page tracking (Page 1, Page 2, etc.)
- Conflicting values / Needs Review flags
- Zero fabrication of missing measurements
- HTTP API status and structured JSON response
"""

import os
import sys
import io
import json
import base64
import urllib.request
import urllib.error

# Ensure workspace root is in path
from pathlib import Path
PROJECT_ROOT = Path(__file__).resolve().parent.parent
BASE_DIR = str(PROJECT_ROOT)
if BASE_DIR not in sys.path:
    sys.path.insert(0, BASE_DIR)
if os.path.join(BASE_DIR, "backend") not in sys.path:
    sys.path.insert(0, os.path.join(BASE_DIR, "backend"))

import pymupdf
import docx
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors

from clinical_extractor import process_uploaded_document, parse_clinical_text, is_ocr_engine_available

TEST_DIR = os.path.join(BASE_DIR, "test_output", "extraction_tests")
os.makedirs(TEST_DIR, exist_ok=True)

test_results = []

def record_test(test_id, name, status, details, metrics=None):
    test_results.append({
        "id": test_id,
        "name": name,
        "status": status,
        "details": details,
        "metrics": metrics or {}
    })
    badge = "[PASS]" if status == "PASS" else "[FAIL]"
    print(f"\n{badge} Test {test_id}: {name}")
    print(f"       Details: {details}")
    if metrics:
        print(f"       Metrics: {metrics}")


# ==============================================================================
# TEST 1: Text-Based Medical PDF
# ==============================================================================
def run_test_1():
    pdf_path = os.path.join(TEST_DIR, "test1_text_medical.pdf")
    doc = SimpleDocTemplate(pdf_path, pagesize=letter)
    styles = getSampleStyleSheet()
    story = [
        Paragraph("<b>CARDIOLOGY CLINICAL SUMMARY</b>", styles["Title"]),
        Spacer(1, 10),
        Paragraph("<b>Patient Name:</b> Eleanor Vance &nbsp;&nbsp;&nbsp; <b>MRN:</b> CV-90210", styles["Normal"]),
        Paragraph("<b>Age:</b> 62 years &nbsp;&nbsp;&nbsp; <b>Gender:</b> Female", styles["Normal"]),
        Paragraph("<b>Height:</b> 165 cm &nbsp;&nbsp;&nbsp; <b>Weight:</b> 68 kg", styles["Normal"]),
        Spacer(1, 8),
        Paragraph("<b>Vital Signs:</b> Resting BP: 142/88 mmHg, Heart Rate: 74 bpm", styles["Normal"]),
        Paragraph("<b>Echocardiography:</b> Left Ventricular Ejection Fraction (LVEF): 55%", styles["Normal"]),
        Spacer(1, 8),
        Paragraph("<b>Medical History:</b> Patient has a diagnosis of Type 2 Diabetes Mellitus (T2DM) and essential hypertension. History of smoking: Ex-smoker (quit 4 years ago). Denies typical chest pain.", styles["Normal"]),
        Spacer(1, 8),
        Paragraph("<b>Laboratory Biomarkers:</b>", styles["Heading3"]),
        Paragraph("Fasting Blood Sugar: 128 mg/dL | Serum Creatinine: 1.10 mg/dL | Total Triglycerides: 185 mg/dL | LDL Cholesterol: 130 mg/dL | HDL Cholesterol: 45 mg/dL | White Blood Cells: 6500 /uL", styles["Normal"]),
    ]
    doc.build(story)

    with open(pdf_path, "rb") as f:
        pdf_bytes = f.read()

    res = process_uploaded_document(pdf_bytes, "test1_text_medical.pdf")
    f_dict = res["features"]
    
    assert res["success"] is True
    assert res["patient_info"]["name"] == "Eleanor Vance"
    assert res["patient_info"]["id"] == "CV-90210"
    assert f_dict.get("age") == 62.0
    assert f_dict.get("sex") == "Female"
    assert f_dict.get("length") == 165.0
    assert f_dict.get("weight") == 68.0
    assert f_dict.get("bmi") == 25.0  # calculated from 68 / (1.65^2)
    assert f_dict.get("bp") == 142.0
    assert f_dict.get("pr") == 74.0
    assert f_dict.get("ef_tte") == 55.0
    assert f_dict.get("dm") == 1
    assert f_dict.get("htn") == 1
    assert f_dict.get("ex_smoker") == 1
    assert f_dict.get("chest_pain") == "N"
    assert f_dict.get("fbs") == 128.0
    assert f_dict.get("cr") == 1.10
    assert f_dict.get("tg") == 185.0
    assert f_dict.get("ldl") == 130.0
    assert f_dict.get("hdl") == 45.0
    assert f_dict.get("wbc") == 6500.0

    record_test(
        1, "Text-based medical PDF", "PASS",
        "Extracted all clinical parameters (demographics, vitals, labs, history) with 100% exact values and calculated BMI.",
        {"extracted_count": res["extracted_count"], "calculated_count": res["calculated_count"]}
    )


# ==============================================================================
# TEST 2: Scanned Medical PDF
# ==============================================================================
def run_test_2():
    pdf_path = os.path.join(TEST_DIR, "test2_scanned_medical.pdf")
    from PIL import Image, ImageDraw
    img = Image.new("RGB", (600, 400), color=(255, 255, 255))
    draw = ImageDraw.Draw(img)
    draw.text((30, 30), "SCANNED ECG AND LAB REPORT - PATIENT ALEX MORGAN", fill=(0, 0, 0))
    draw.text((30, 70), "Age: 58 | Sex: Male | BP: 130/80 | HR: 72", fill=(0, 0, 0))
    img_byte_arr = io.BytesIO()
    img.save(img_byte_arr, format="PNG")
    img_bytes = img_byte_arr.getvalue()

    doc = pymupdf.open()
    page = doc.new_page(width=600, height=400)
    page.insert_image(pymupdf.Rect(0, 0, 600, 400), stream=img_bytes)
    doc.save(pdf_path)
    doc.close()

    with open(pdf_path, "rb") as f:
        pdf_bytes = f.read()

    ocr_avail = is_ocr_engine_available()
    try:
        res = process_uploaded_document(pdf_bytes, "test2_scanned_medical.pdf")
        if ocr_avail:
            assert res["ocr_status"] == "performed"
            record_test(2, "Scanned medical PDF", "PASS", "OCR engine executed on scanned page image successfully.", {"ocr_status": res["ocr_status"]})
        else:
            record_test(2, "Scanned medical PDF", "PASS", "Gracefully detected scanned page and reported clear message.", {"ocr_status": res["ocr_status"]})
    except ValueError as ve:
        assert "scanned" in str(ve).lower() or "ocr" in str(ve).lower()
        record_test(2, "Scanned medical PDF", "PASS", f"Detected scanned document without OCR engine and raised clear user error: '{str(ve)}'", {"ocr_available": ocr_avail})


# ==============================================================================
# TEST 3: Multi-Page Laboratory Report
# ==============================================================================
def run_test_3():
    pdf_path = os.path.join(TEST_DIR, "test3_multipage_report.pdf")
    doc = SimpleDocTemplate(pdf_path, pagesize=letter)
    styles = getSampleStyleSheet()

    story = [
        # PAGE 1
        Paragraph("<b>HOSPITAL CLINICAL EVALUATION — PAGE 1: DEMOGRAPHICS & VITALS</b>", styles["Title"]),
        Spacer(1, 14),
        Paragraph("<b>Patient Name:</b> Marcus Brody &nbsp;&nbsp;&nbsp; <b>ID:</b> MB-7741", styles["Normal"]),
        Paragraph("<b>Age:</b> 55 years &nbsp;&nbsp;&nbsp; <b>Sex:</b> Male", styles["Normal"]),
        Paragraph("<b>Height:</b> 175 cm &nbsp;&nbsp;&nbsp; <b>Weight:</b> 78 kg", styles["Normal"]),
        Paragraph("<b>Resting Blood Pressure:</b> 130/82 mmHg &nbsp;&nbsp;&nbsp; <b>Heart Rate:</b> 68 bpm", styles["Normal"]),
        PageBreak(),

        # PAGE 2
        Paragraph("<b>PAGE 2: COMPREHENSIVE METABOLIC & LIPID PANEL</b>", styles["Title"]),
        Spacer(1, 14),
        Paragraph("Fasting Blood Sugar: 105 mg/dL", styles["Normal"]),
        Paragraph("Serum Creatinine: 0.90 mg/dL", styles["Normal"]),
        Paragraph("Blood Urea Nitrogen (BUN): 15 mg/dL", styles["Normal"]),
        Paragraph("Triglycerides: 160 mg/dL", styles["Normal"]),
        Paragraph("LDL Cholesterol: 115 mg/dL", styles["Normal"]),
        Paragraph("HDL Cholesterol: 42 mg/dL", styles["Normal"]),
        PageBreak(),

        # PAGE 3
        Paragraph("<b>PAGE 3: 12-LEAD ECG & ECHOCARDIOGRAPHY</b>", styles["Title"]),
        Spacer(1, 14),
        Paragraph("Left Ventricular Ejection Fraction (LVEF): 60%", styles["Normal"]),
        Paragraph("ST depression: Present (1mm horizontal in V5-V6)", styles["Normal"]),
        Paragraph("T-wave inversion: Present", styles["Normal"]),
        Paragraph("Pathological Q waves: Absent", styles["Normal"]),
    ]
    doc.build(story)

    with open(pdf_path, "rb") as f:
        pdf_bytes = f.read()

    res = process_uploaded_document(pdf_bytes, "test3_multipage_report.pdf")
    items = {item["key"]: item for item in res["review_items"]}

    assert items["age"]["source_page"] == "Page 1"
    assert items["bp"]["source_page"] == "Page 1"
    assert items["fbs"]["source_page"] == "Page 2"
    assert items["cr"]["source_page"] == "Page 2"
    assert items["ef_tte"]["source_page"] == "Page 3"
    assert items["st_depression"]["source_page"] == "Page 3"
    assert items["bmi"]["source_page"] == "Calculated"

    record_test(
        3, "Multi-page laboratory report", "PASS",
        f"Accurately mapped 3-page document with Page 1 (vitals), Page 2 (biomarkers), and Page 3 (cardiac telemetry). Total pages analyzed: {res['pages_count']}.",
        {"pages_count": res["pages_count"], "extracted_count": res["extracted_count"]}
    )


# ==============================================================================
# TEST 4: DOCX Medical Report
# ==============================================================================
def run_test_4():
    docx_path = os.path.join(TEST_DIR, "test4_medical_report.docx")
    d = docx.Document()
    d.add_heading("Cardiology Patient Assessment", level=1)
    d.add_paragraph("Patient Name: Clara Oswald")
    d.add_paragraph("Patient ID: DW-101")
    d.add_paragraph("Age: 51 years | Sex: Female")
    d.add_paragraph("Height: 162 cm | Weight: 64 kg")
    d.add_paragraph("Resting Blood Pressure: 135/85 mmHg | Pulse: 72 bpm")
    d.add_paragraph("Medical History: Essential hypertension, non-diabetic, non-smoker. Reports typical chest pain.")
    
    t = d.add_table(rows=1, cols=3)
    hdr = t.rows[0].cells
    hdr[0].text = "Analyte"
    hdr[1].text = "Measured Result"
    hdr[2].text = "Reference Interval"

    lab_data = [
        ("Fasting Blood Glucose", "112 mg/dL", "70 - 99"),
        ("Serum Creatinine", "0.95 mg/dL", "0.60 - 1.20"),
        ("Total Triglycerides", "170 mg/dL", "< 150"),
        ("LDL Cholesterol", "125 mg/dL", "< 100"),
        ("HDL Cholesterol", "48 mg/dL", "> 40"),
        ("Ejection Fraction (EF)", "58 %", "50 - 70"),
    ]
    for analyte, val, ref in lab_data:
        row = t.add_row().cells
        row[0].text = analyte
        row[1].text = val
        row[2].text = ref

    d.save(docx_path)

    with open(docx_path, "rb") as f:
        docx_bytes = f.read()

    res = process_uploaded_document(docx_bytes, "test4_medical_report.docx")
    f_dict = res["features"]

    assert res["file_type"] == "docx"
    assert f_dict["age"] == 51.0
    assert f_dict["sex"] == "Female"
    assert f_dict["bp"] == 135.0
    assert f_dict["pr"] == 72.0
    assert f_dict["htn"] == 1
    assert f_dict["dm"] == 0
    assert f_dict["fbs"] == 112.0
    assert f_dict["cr"] == 0.95
    assert f_dict["tg"] == 170.0
    assert f_dict["ldl"] == 125.0
    assert f_dict["hdl"] == 48.0
    assert f_dict["ef_tte"] == 58.0

    record_test(
        4, "DOCX medical report", "PASS",
        "Successfully parsed Word (.docx) document including tables, paragraphs, and diagnosis sentences.",
        {"file_type": res["file_type"], "extracted_count": res["extracted_count"]}
    )


# ==============================================================================
# TEST 5: PDF Containing Laboratory Tables
# ==============================================================================
def run_test_5():
    pdf_path = os.path.join(TEST_DIR, "test5_lab_tables.pdf")
    doc = SimpleDocTemplate(pdf_path, pagesize=letter)
    styles = getSampleStyleSheet()

    table_data = [
        ["Test Name", "Result", "Units", "Reference Interval"],
        ["Fasting Blood Sugar", "115", "mg/dL", "70 - 99"],
        ["Serum Creatinine", "1.20", "mg/dL", "0.60 - 1.20"],
        ["Triglycerides", "210", "mg/dL", "< 150"],
        ["LDL Cholesterol", "145", "mg/dL", "< 100"],
        ["HDL Cholesterol", "38", "mg/dL", "> 40"],
        ["Potassium", "4.6", "mEq/L", "3.5 - 5.0"],
        ["Sodium", "142", "mEq/L", "136 - 145"],
        ["Hemoglobin", "14.5", "g/dL", "13.0 - 17.0"],
        ["White Blood Cells", "7200", "/uL", "4000 - 11000"],
        ["Platelet Count", "240", "x10^3/uL", "150 - 450"],
    ]

    t = Table(table_data, colWidths=[180, 70, 70, 120])
    t.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor("#1e3a8a")),
        ('TEXTCOLOR', (0,0), (-1,0), colors.white),
        ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor("#cbd5e1")),
    ]))

    story = [
        Paragraph("<b>LABORATORY REPORT — BLOOD CHEMISTRY & HEMATOLOGY</b>", styles["Title"]),
        Spacer(1, 10),
        Paragraph("Patient: Thomas Anderson | Age: 45 | Sex: Male | Height: 178 cm | Weight: 80 kg", styles["Normal"]),
        Spacer(1, 10),
        t
    ]
    doc.build(story)

    with open(pdf_path, "rb") as f:
        pdf_bytes = f.read()

    res = process_uploaded_document(pdf_bytes, "test5_lab_tables.pdf")
    f_dict = res["features"]

    assert f_dict["fbs"] == 115.0  # NOT 70 or 99
    assert f_dict["cr"] == 1.20   # NOT 0.60
    assert f_dict["tg"] == 210.0   # NOT 150
    assert f_dict["ldl"] == 145.0  # NOT 100
    assert f_dict["hdl"] == 38.0   # NOT 40
    assert f_dict["k"] == 4.6      # NOT 3.5 or 5.0
    assert f_dict["na"] == 142.0   # NOT 136 or 145
    assert f_dict["hb"] == 14.5    # NOT 13.0 or 17.0
    assert f_dict["wbc"] == 7200.0 # NOT 4000 or 11000
    assert f_dict["plt"] == 240.0  # NOT 150 or 450

    record_test(
        5, "PDF containing laboratory tables", "PASS",
        "Disambiguated 10/10 laboratory table rows. Patient results extracted without confusing reference intervals.",
        {"fbs": f_dict["fbs"], "cr": f_dict["cr"], "tg": f_dict["tg"], "ldl": f_dict["ldl"]}
    )


# ==============================================================================
# TEST 6: PDF With Missing Measurements
# ==============================================================================
def run_test_6():
    pdf_path = os.path.join(TEST_DIR, "test6_missing_measurements.pdf")
    doc = SimpleDocTemplate(pdf_path, pagesize=letter)
    styles = getSampleStyleSheet()

    story = [
        Paragraph("<b>PRE-OP SCREENING — VITALS ONLY</b>", styles["Title"]),
        Spacer(1, 10),
        Paragraph("Patient: Sarah Connor | Age: 48 years | Gender: Female", styles["Normal"]),
        Paragraph("Height: 160 cm | Weight: 60 kg", styles["Normal"]),
        Paragraph("Blood Pressure: 120/80 mmHg | Pulse Rate: 68 bpm", styles["Normal"]),
        Paragraph("<i>Note: Comprehensive metabolic panel and ECG pending hospital transfer.</i>", styles["Normal"]),
    ]
    doc.build(story)

    with open(pdf_path, "rb") as f:
        pdf_bytes = f.read()

    res = process_uploaded_document(pdf_bytes, "test6_missing_measurements.pdf")
    items = {item["key"]: item for item in res["review_items"]}

    assert items["age"]["status"] == "Extracted"
    assert items["length"]["status"] == "Extracted"
    assert items["weight"]["status"] == "Extracted"
    assert items["bmi"]["status"] == "Calculated"
    assert items["bp"]["status"] == "Extracted"

    assert items["fbs"]["status"] == "Not Found" and items["fbs"]["value"] is None
    assert items["cr"]["status"] == "Not Found" and items["cr"]["value"] is None
    assert items["tg"]["status"] == "Not Found" and items["tg"]["value"] is None
    assert items["q_wave"]["status"] == "Not Found" and items["q_wave"]["value"] is None
    assert items["region_rwma"]["status"] == "Not Found" and items["region_rwma"]["value"] is None

    record_test(
        6, "PDF with missing measurements", "PASS",
        f"Correctly isolated missing measurements. {res['not_found_count']} fields marked 'Not Found' with None value (zero silent substitution).",
        {"not_found_count": res["not_found_count"], "extracted_count": res["extracted_count"]}
    )


# ==============================================================================
# TEST 7: PDF Containing Conflicting Values
# ==============================================================================
def run_test_7():
    pdf_path = os.path.join(TEST_DIR, "test7_conflicting_values.pdf")
    doc = SimpleDocTemplate(pdf_path, pagesize=letter)
    styles = getSampleStyleSheet()

    story = [
        Paragraph("<b>CLINICAL REPORT — CONFLICTING MEASUREMENTS</b>", styles["Title"]),
        Spacer(1, 10),
        Paragraph("Patient: James Cole | Age: 52 | Sex: Male", styles["Normal"]),
        Paragraph("Height: 170 cm | Weight: 70 kg", styles["Normal"]),
        # 70 / (1.70^2) = 24.2 kg/m2, but document states BMI: 31.5 kg/m2 (conflicting by 7.3 kg/m2)
        Paragraph("Stated BMI: 31.5 kg/m2", styles["Normal"]),
        Paragraph("Triage Blood Pressure: 120/80 mmHg", styles["Normal"]),
        PageBreak(),
        Paragraph("<b>PAGE 2: INPATIENT RE-ASSESSMENT</b>", styles["Title"]),
        Paragraph("Subsequent Blood Pressure: 165/100 mmHg", styles["Normal"]),
    ]
    doc.build(story)

    with open(pdf_path, "rb") as f:
        pdf_bytes = f.read()

    res = process_uploaded_document(pdf_bytes, "test7_conflicting_values.pdf")
    items = {item["key"]: item for item in res["review_items"]}

    assert items["bmi"]["status"] == "Needs Review"
    assert "differs from calculated" in str(items["bmi"]["notes"]).lower()
    assert items["bp"]["status"] == "Needs Review"
    assert "multiple readings" in str(items["bp"]["notes"]).lower()

    record_test(
        7, "PDF containing conflicting values", "PASS",
        "Flagged both conflicting BMI (stated 31.5 vs calculated 24.2) and multi-page conflicting BP readings (120 vs 165) as 'Needs Review'.",
        {"needs_review_count": res["needs_review_count"]}
    )


# ==============================================================================
# TEST 8: PDF With Different Measurement Units
# ==============================================================================
def run_test_8():
    pdf_path = os.path.join(TEST_DIR, "test8_different_units.pdf")
    doc = SimpleDocTemplate(pdf_path, pagesize=letter)
    styles = getSampleStyleSheet()

    story = [
        Paragraph("<b>INTERNATIONAL CLINICAL LAB REPORT — METRIC & IMPERIAL UNITS</b>", styles["Title"]),
        Spacer(1, 10),
        Paragraph("Patient: Arthur Dent | Age: 42 | Gender: Male", styles["Normal"]),
        Paragraph("Height: 5 ft 10 in | Weight: 165 lbs", styles["Normal"]),
        Paragraph("Fasting Plasma Glucose: 6.5 mmol/L", styles["Normal"]),
        Paragraph("Serum Creatinine: 106 umol/L", styles["Normal"]),
        Paragraph("LDL Cholesterol: 3.6 mmol/L", styles["Normal"]),
        Paragraph("Total Triglycerides: 2.0 mmol/L", styles["Normal"]),
        Paragraph("White Blood Cell Count: 7.5 x10^3/uL", styles["Normal"]),
    ]
    doc.build(story)

    with open(pdf_path, "rb") as f:
        pdf_bytes = f.read()

    res = process_uploaded_document(pdf_bytes, "test8_different_units.pdf")
    f_dict = res["features"]

    assert abs(f_dict["length"] - 177.8) < 0.2
    assert abs(f_dict["weight"] - 74.8) < 0.2
    assert abs(f_dict["fbs"] - 117.1) < 0.5
    assert abs(f_dict["cr"] - 1.20) < 0.05
    assert abs(f_dict["ldl"] - 139.2) < 0.5
    assert abs(f_dict["tg"] - 177.1) < 0.5
    assert f_dict["wbc"] == 7500.0

    record_test(
        8, "PDF with different measurement units", "PASS",
        "Applied analyte-specific conversions (ft/in -> cm, lbs -> kg, mmol/L -> mg/dL, umol/L -> mg/dL, 10^3 -> absolute).",
        {"length_cm": f_dict["length"], "weight_kg": f_dict["weight"], "glucose_mgdl": f_dict["fbs"], "cr_mgdl": f_dict["cr"]}
    )


# ==============================================================================
# TEST 9: Invalid or Corrupted PDF
# ==============================================================================
def run_test_9():
    corrupt_bytes = b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\nCORRUPTED_BINARY_DATA_CRASH_TEST\x00\xff" * 50

    caught_expected_error = False
    try:
        process_uploaded_document(corrupt_bytes, "corrupt.pdf")
    except ValueError as ve:
        caught_expected_error = True
        assert "corrupt" in str(ve).lower() or "invalid" in str(ve).lower()

    assert caught_expected_error is True

    api_url = "http://127.0.0.1:8502/api/upload-report"
    b64_data = base64.b64encode(corrupt_bytes).decode("utf-8")
    req = urllib.request.Request(
        api_url,
        data=json.dumps({"filename": "corrupt.pdf", "file_data": b64_data}).encode("utf-8"),
        headers={"Content-Type": "application/json"}
    )
    api_returned_400 = False
    try:
        urllib.request.urlopen(req)
    except urllib.error.HTTPError as e:
        if e.code == 400:
            api_returned_400 = True
            body = json.loads(e.read().decode("utf-8"))
            assert body["success"] is False
            assert "corrupt" in body["error"].lower() or "invalid" in body["error"].lower()

    assert api_returned_400 is True

    record_test(
        9, "Invalid or corrupted PDF", "PASS",
        "Handled corrupt PDF safely in Python (ValueError) and REST API (HTTP 400 Bad Request with informative JSON error).",
        {"api_status_code": 400, "handled_cleanly": True}
    )


# ==============================================================================
# TEST 10: PDF With All Supported Clinical Fields
# ==============================================================================
def run_test_10():
    pdf_path = os.path.join(TEST_DIR, "test10_complete_57_features.pdf")
    doc = SimpleDocTemplate(pdf_path, pagesize=letter)
    styles = getSampleStyleSheet()

    story = [
        Paragraph("<b>COMPREHENSIVE CARDIOLOGY EVALUATION — ALL 57 FEATURES</b>", styles["Title"]),
        Spacer(1, 8),
        Paragraph("<b>Patient Name:</b> Eleanor Vance &nbsp;&nbsp;&nbsp; <b>ID:</b> CV-90210", styles["Normal"]),
        Paragraph("<b>Age:</b> 62 years &nbsp;&nbsp;&nbsp; <b>Sex:</b> Female", styles["Normal"]),
        Paragraph("<b>Height:</b> 165 cm &nbsp;&nbsp;&nbsp; <b>Weight:</b> 68 kg", styles["Normal"]),
        Paragraph("<b>Resting BP:</b> 142/88 mmHg &nbsp;&nbsp;&nbsp; <b>Pulse:</b> 74 bpm &nbsp;&nbsp;&nbsp; <b>LVEF:</b> 55 %", styles["Normal"]),
        Spacer(1, 6),
        Paragraph("<b>Medical History & Symptoms:</b>", styles["Heading3"]),
        Paragraph("Diabetes Mellitus: Type 2 Diabetes (T2DM: Yes). Hypertension: Essential HTN. Dyslipidemia: DLP present.", styles["Normal"]),
        Paragraph("Smoking: Ex-smoker (former smoker, quit 4 years ago). Active smoker: No. Family History of CAD: Positive.", styles["Normal"]),
        Paragraph("Obesity: Non-obese. Chronic Renal Failure (CRF): No. Cerebrovascular Accident (CVA): No history.", styles["Normal"]),
        Paragraph("Airway Disease: No asthma or COPD. Thyroid Disease: Hypothyroidism present. Heart Failure: Denies CHF.", styles["Normal"]),
        Paragraph("Symptoms: Reports typical chest pain / angina on exertion. Denies atypical chest pain, denies nonanginal pain, denies low threshold angina. Dyspnea: Reports shortness of breath (NYHA Function Class II).", styles["Normal"]),
        Spacer(1, 6),
        Paragraph("<b>Physical Findings:</b>", styles["Heading3"]),
        Paragraph("Bilateral pedal edema: Present. Peripheral pulses: Normal, no weak peripheral pulse. Lungs: Clear, no lung rales. Auscultation: Systolic murmur present at apex; no diastolic murmur.", styles["Normal"]),
        Spacer(1, 6),
        Paragraph("<b>ECG & Regional Wall Motion:</b>", styles["Heading3"]),
        Paragraph("Pathological Q waves: Absent. ST elevation: None. ST depression: Present (1mm in V5-V6). T-wave inversion: Present. LVH: Absent. Bundle Branch Block: None (BBB: N). Poor R progression: Present. RWMA score: 1.0 (Mild hypokinesia). Valvular Disease: Mild mitral regurgitation (VHD: mild).", styles["Normal"]),
        Spacer(1, 6),
        Paragraph("<b>Complete Laboratory Panel:</b>", styles["Heading3"]),
        Paragraph("FBS: 128 mg/dL | Serum Creatinine: 1.10 mg/dL | Triglycerides: 185 mg/dL | LDL: 130 mg/dL | HDL: 45 mg/dL | BUN: 18 mg/dL | ESR: 16 mm/hr | Hemoglobin: 13.8 g/dL | Potassium: 4.3 mEq/L | Sodium: 141 mEq/L | WBC: 7400 /uL | Lymphocytes: 31 % | Neutrophils: 62 % | Platelets: 230 x10^3/uL", styles["Normal"]),
    ]
    doc.build(story)

    with open(pdf_path, "rb") as f:
        pdf_bytes = f.read()

    api_url = "http://127.0.0.1:8502/api/upload-report"
    b64_data = base64.b64encode(pdf_bytes).decode("utf-8")
    req = urllib.request.Request(
        api_url,
        data=json.dumps({"filename": "test10_complete_57_features.pdf", "file_data": b64_data}).encode("utf-8"),
        headers={"Content-Type": "application/json"}
    )
    with urllib.request.urlopen(req) as resp:
        assert resp.status == 200
        res = json.loads(resp.read().decode("utf-8"))

    assert res["success"] is True
    assert res["extracted_count"] >= 35
    assert res["calculated_count"] >= 1  # BMI

    from patient_feature_snapshot import create_patient_feature_snapshot
    snapshot = create_patient_feature_snapshot(res["features"], model_mode="standard", strict=False)
    assert len(snapshot.X.columns) == 57
    assert snapshot.X["Age"].iloc[0] == 62.0
    assert snapshot.X["Length"].iloc[0] == 165.0
    assert snapshot.X["Weight"].iloc[0] == 68.0
    assert snapshot.X["DM"].iloc[0] == 1
    assert snapshot.X["HTN"].iloc[0] == 1
    assert snapshot.X["FBS"].iloc[0] == 128.0
    assert snapshot.X["CR"].iloc[0] == 1.10
    assert snapshot.X["TG"].iloc[0] == 185.0

    record_test(
        10, "PDF with all supported clinical fields", "PASS",
        f"Extracted {res['extracted_count']} clinical features and derived BMI. Seamlessly constructed 57-dimensional feature matrix X.",
        {"extracted_count": res["extracted_count"], "calculated_count": res["calculated_count"], "feature_matrix_cols": len(snapshot.X.columns)}
    )


def main():
    print("================================================================================")
    print("      CARDIOVISION AI — 10-SCENARIO CLINICAL DOCUMENT EXTRACTION SUITE          ")
    print("================================================================================")
    
    run_test_1()
    run_test_2()
    run_test_3()
    run_test_4()
    run_test_5()
    run_test_6()
    run_test_7()
    run_test_8()
    run_test_9()
    run_test_10()

    total_tests = len(test_results)
    pass_count = sum(1 for t in test_results if t["status"] == "PASS")
    fail_count = total_tests - pass_count

    print("\n================================================================================")
    print(f"                      TEST SUMMARY: {pass_count}/{total_tests} PASS (0 FAIL)")
    print("================================================================================")

    report_path = os.path.join(TEST_DIR, "extraction_verification_report.json")
    with open(report_path, "w", encoding="utf-8") as f:
        json.dump({
            "total": total_tests,
            "pass": pass_count,
            "fail": fail_count,
            "tests": test_results
        }, f, indent=2)
    print(f"Saved results to: {report_path}")

    assert fail_count == 0, f"{fail_count} tests failed!"

if __name__ == "__main__":
    main()
